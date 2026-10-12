#!/usr/bin/env python3
"""Build a formula-driven cloud cost projection workbook (+ CSV) from a JSON config.

Usage:  python3 build_model.py <config.json> [output_dir]
Produces <output_dir>/<company> Cloud Cost Projection <first>-<last>.xlsx and .csv, recalculates all
formulas with LibreOffice when the xlsx skill's recalc.py is available, and prints the tie-outs.
See references/config.md for every field. All money in USD unless the config says otherwise;
the report currency (EUR by default) is only used for the KPI ratio and the cap.
"""
import json, os, subprocess, sys, csv
from datetime import date
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

CFG = json.load(open(sys.argv[1]))
OUTDIR = sys.argv[2] if len(sys.argv) > 2 else "."
os.makedirs(OUTDIR, exist_ok=True)

# ------------------------------------------------------------------ months
y, m = map(int, CFG["months"]["start"].split("-"))
MONTHS, YEARS = [], []
for k in range(CFG["months"]["count"]):
    mm = (m - 1 + k) % 12 + 1; yy = y + (m - 1 + k) // 12
    MONTHS.append(date(yy, mm, 1).strftime("%b %Y")); YEARS.append(yy)
N = len(MONTHS); PLAN = CFG["months"]["plan_year"]
PLAN_IDX = [j for j in range(N) if YEARS[j] == PLAN]           # 0-based month indexes in the plan year
P0 = PLAN_IDX[0]                                                # first plan-year month (0-based)
SC = [("P", "Pessimistic"), ("R", "Realistic"), ("O", "Optimistic")]
CCY = CFG.get("report_currency", "EUR")
COMPANY = CFG.get("company", "Company")
OUT = os.path.join(OUTDIR, f"{COMPANY} Cloud Cost Projection {MONTHS[0].replace(' ', '')}-{MONTHS[-1].replace(' ', '')}.xlsx")

# ------------------------------------------------------------------ styles
F = "Arial"
f_b = Font(name=F, size=10); f_bold = Font(name=F, size=10, bold=True); f_h = Font(name=F, size=10, bold=True, color="FFFFFF")
f_t = Font(name=F, size=13, bold=True); f_in = Font(name=F, size=10, color="0000FF"); f_note = Font(name=F, size=9, italic=True, color="595959")
fill_h = PatternFill("solid", fgColor="1A1A1A"); fill_in = PatternFill("solid", fgColor="FFF2CC"); fill_sub = PatternFill("solid", fgColor="F2F2F2")
fill_com = PatternFill("solid", fgColor="FCE4D6"); fill_x = PatternFill("solid", fgColor="E7E6E6"); wrap = Alignment(wrap_text=True, vertical="top")
USD = '#,##0;(#,##0);-'; USD2 = '#,##0.00;(#,##0.00);-'; PCT = '0.0%'
def hdr(ws, row, labels, widths=None):
    for c, v in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=c, value=v); cell.font = f_h; cell.fill = fill_h; cell.alignment = wrap
    if widths:
        for c, w in enumerate(widths, start=1): ws.column_dimensions[L(c)].width = w
def blue(cell, fmt=None):
    cell.font = f_in; cell.fill = fill_in
    if fmt: cell.number_format = fmt

PRODUCTS = [p["name"] for p in CFG["products"]]; ELAST = {p["name"]: p["elasticity"] for p in CFG["products"]}
PROJECTS = CFG["projects"]; INSCOPE = [p for p in PROJECTS if p["in_scope"] == 1]
ITEMS = CFG["items"]; AWS = CFG["aws_accounts"]; INIT = CFG["initiatives"]; TERM = CFG.get("terminating_commitments", [])
ENV = CFG.get("envelope"); SCEN = CFG["scenarios"]; SW = CFG.get("switches", {})
FLOOR_ITEM = next((it for it in ITEMS if it.get("floor")), None)
SUPPORT_ITEM = next((it for it in ITEMS if it.get("support_ratio")), None)

wb = Workbook()
# ================================================================ Inputs
ws = wb.active; ws.title = "Inputs"
ws["A1"] = f"{COMPANY}: cloud cost projection by project and account, {MONTHS[0]} to {MONTHS[-1]}. Blue = editable. USD unless marked {CCY}."; ws["A1"].font = f_t
hdr(ws, 3, ["Parameter", "Value", "Source / note"], [62, 18, 110])
single = [
 (f"FX USD per {CCY}", CFG["fx_usd_per_report_ccy"], "Finance planning rate"),
 (f"Current-year revenue ({CCY})", CFG["revenue"]["current_year"], CFG["revenue"].get("note", "")),
 (f"Plan-year revenue ({CCY})", CFG["revenue"]["plan_year"], "Check the revenue basis matches the cost scope (same entities in and out)"),
 ("Cloud cost cap (share of revenue, KPI)", CFG["kpi_cap_share"], "Company KPI"),
 ("Support ratio (share of GCP project usage)", f"={SUPPORT_ITEM['aug_invoice']}/{CFG['reconciliation']['project_usage']}" if SUPPORT_ITEM else 0, "Baseline invoice: support charge / project usage"),
 ("Floor vendor: contract monthly minimum", FLOOR_ITEM["floor"]["min_monthly"] if FLOOR_ITEM else 0, FLOOR_ITEM["floor"].get("note", "") if FLOOR_ITEM else ""),
 ("Floor vendor: uplift on fees (e.g. premium support share)", FLOOR_ITEM["floor"].get("uplift_share", 0) if FLOOR_ITEM else 0, ""),
 ("Payment lag for usage-billed lines (months)", CFG["payment_lag_months"], "Reseller invoices month M early M+1 at Net 30: lag 2"),
 ("Excluded from every total", ", ".join(CFG.get("excluded_memo", [])) or "none", "e.g. Workspace / IT items on the same invoice"),
 ("Switch: floor-vendor plan-year fee as single invoice (1) or spread monthly (0)", SW.get("floor_year_fee_single", 1), "Cash events tagged with this switch follow it"),
 ("Run-rate baseline", CFG["baseline_label"], "Project x product matrix on the Projects sheet reconciles to the invoice sub-total"),
 ("Terminating commitment: last month billed (t index, 1 = first month; 0 = ended)", TERM[0]["last_month_index"] if TERM else 0, TERM[0].get("note", "") if TERM else ""),
 ("Allocation key for account-level items", "GCP usage share", "Account-level vendors are spread over every project by its share of baseline usage; excluded projects carry their share out"),
 ("Cash flow basis: 1 = whole account (what the company pays), 0 = in-scope share", 1, "Reseller invoices are paid in full regardless of scope"),
 (f"Innovation envelope, plan year ({CCY})", ENV["amount_report_ccy"] if ENV else 0, "Discretionary, company-level line; not allocated to projects"),
 ("Envelope: share billed through the cloud reseller (rest direct-billed, paid next month)", ENV["reseller_billed_share"] if ENV else 0, ""),
 ("Envelope: count inside the KPI (1 = yes, 0 = beside it)", ENV["inside_kpi"] if ENV else 1, "Both ratios are shown"),
]
for i, (k, v, s) in enumerate(single, start=4):
    ws.cell(row=i, column=1, value=k).font = f_b
    c = ws.cell(row=i, column=2, value=v); blue(c)
    ws.cell(row=i, column=3, value=s).font = f_note; ws.cell(row=i, column=3).alignment = wrap
for cell, fmt in (("B5", USD), ("B6", USD), ("B7", PCT), ("B8", PCT), ("B9", USD), ("B10", PCT), ("B18", USD), ("B19", PCT)): ws[cell].number_format = fmt
FX = "Inputs!$B$4"; REV_CUR = "Inputs!$B$5"; REV_PLAN = "Inputs!$B$6"; CAP = "Inputs!$B$7"; SUPP = "Inputs!$B$8"; FMIN = "Inputs!$B$9"; FUP = "Inputs!$B$10"
LAG = "Inputs!$B$11"; SWITCH = "Inputs!$B$13"; TERMEND = "Inputs!$B$15"; CASHB = "Inputs!$B$17"; ENVAMT = "Inputs!$B$18"; ENVRES = "Inputs!$B$19"; ENVKPI = "Inputs!$B$20"
SR = 23  # scenario table header row
hdr(ws, SR, ["Scenario parameter", "Pessimistic", "Realistic", "Optimistic", "Note"])
scen = [
 ("Volume growth, plan year (YoY)", SCEN["growth_yoy"], "Drives every line through its elasticity"),
 ("Monthly volume growth (derived)", None, ""),
 ("Next milestone credit: month index it posts (1 = first month)", SCEN["credit_month_index"], SCEN.get("credit_note", "")),
 ("Milestone credit amount (USD, whole billing account)", SCEN["credit_amount"], ""),
 ("Envelope utilization (share actually spent)", (ENV or {}).get("utilization", {"P": 1, "R": 1, "O": 1}), "A budget is a cap; scenarios vary the spend against it"),
]
for i, (k, v, s) in enumerate(scen, start=SR + 1):
    ws.cell(row=i, column=1, value=k).font = f_b
    for col, key in ((2, "P"), (3, "R"), (4, "O")):
        val = f"=(1+{L(col)}{SR+1})^(1/12)-1" if v is None else (v[key] if isinstance(v, dict) else v)
        c = ws.cell(row=i, column=col, value=val); blue(c, PCT if i in (SR + 1, SR + 2, SR + 5) else USD)
    ws.cell(row=i, column=5, value=s).font = f_note; ws.cell(row=i, column=5).alignment = wrap
ws.column_dimensions["D"].width = 14; ws.column_dimensions["E"].width = 100
GROW = {k: f"Inputs!${c}${SR+2}" for k, c in (("P", "B"), ("R", "C"), ("O", "D"))}
CREDM = {k: f"Inputs!${c}${SR+3}" for k, c in (("P", "B"), ("R", "C"), ("O", "D"))}
CREDA = {k: f"Inputs!${c}${SR+4}" for k, c in (("P", "B"), ("R", "C"), ("O", "D"))}
ENVUTIL = {k: f"Inputs!${c}${SR+5}" for k, c in (("P", "B"), ("R", "C"), ("O", "D"))}

# ================================================================ Projects (baseline matrix + allocation)
ws = wb.create_sheet("Projects baseline")
ws["A1"] = f"{CFG['baseline_label']} invoice by project and product line (USD before VAT). In scope = 1 enters the projection. Account-level items are allocated to every project by its usage share (column O)."; ws["A1"].font = f_note
item_names = [it["name"] for it in ITEMS]; NPR = len(PRODUCTS)
cols = ["Project", "Group", "Env", "In scope"] + PRODUCTS + ["Usage total", "Usage share (all projects)"] + [f"Allocated: {n}" for n in item_names] + ["All-in"]
hdr(ws, 2, cols, [34, 12, 8, 8] + [16] * NPR + [16, 12] + [16] * len(ITEMS) + [18])
r0 = 3; NP = len(PROJECTS); last = r0 + NP - 1
COL_TOTAL = 4 + NPR + 1; COL_SHARE = COL_TOTAL + 1; COL_ITEM0 = COL_SHARE + 1; COL_ALLIN = COL_ITEM0 + len(ITEMS)
for i, p in enumerate(PROJECTS):
    r = r0 + i
    for c, v in enumerate([p["id"], p["group"], p["env"], p["in_scope"]], start=1):
        cell = ws.cell(row=r, column=c, value=v); cell.font = f_b
        if c == 4: blue(cell)
    for k, v in enumerate(p["values"]):
        blue(ws.cell(row=r, column=5 + k, value=v), USD2)
    ws.cell(row=r, column=COL_TOTAL, value=f"=SUM({L(5)}{r}:{L(4+NPR)}{r})").number_format = USD2
    ws.cell(row=r, column=COL_SHARE, value=f"={L(COL_TOTAL)}{r}/{L(COL_TOTAL)}${last+1}").number_format = '0.00%'
    for k, it in enumerate(ITEMS):
        val = f"={SUPP}*{L(COL_TOTAL)}{r}" if it.get("support_ratio") else f"={it['run_rate']}*{L(COL_SHARE)}{r}"
        ws.cell(row=r, column=COL_ITEM0 + k, value=val).number_format = USD
    ws.cell(row=r, column=COL_ALLIN, value=f"={L(COL_TOTAL)}{r}+SUM({L(COL_ITEM0)}{r}:{L(COL_ITEM0+len(ITEMS)-1)}{r})").number_format = USD
    if p["in_scope"] == 0:
        for c in range(1, COL_ALLIN + 1): ws.cell(row=r, column=c).fill = fill_x
rt = last + 1
for lbl, rr in (("All projects", rt), ("In scope", rt + 1), ("Excluded", rt + 2)): ws.cell(row=rr, column=1, value=lbl).font = f_bold
for c in list(range(5, COL_TOTAL + 1)) + list(range(COL_ITEM0, COL_ALLIN + 1)):
    col = L(c)
    ws.cell(row=rt, column=c, value=f"=SUM({col}{r0}:{col}{last})").number_format = USD
    ws.cell(row=rt + 1, column=c, value=f"=SUMPRODUCT($D${r0}:$D${last},{col}{r0}:{col}{last})").number_format = USD
    ws.cell(row=rt + 2, column=c, value=f"={col}{rt}-{col}{rt+1}").number_format = USD
ws.cell(row=rt + 1, column=COL_SHARE, value=f"=SUMPRODUCT($D${r0}:$D${last},{L(COL_SHARE)}{r0}:{L(COL_SHARE)}{last})").number_format = '0.00%'
REC = CFG["reconciliation"]
ws.cell(row=rt + 4, column=1, value="Reconciliation to the baseline invoice").font = f_bold
rec = [("Project usage (all projects)", f"={L(COL_TOTAL)}{rt}", f"{REC['project_usage']:,} in the invoice"),
       ("Account-level block (support, floor vendor, offsets)", REC["account_block"], ""), ("Marketplace projects (data platform vendors)", REC["marketplace"], ""),
       ("Excluded memo items (IT)", REC.get("excluded_items", 0), ""),
       ("Sum", f"=B{rt+5}+B{rt+6}+B{rt+7}+B{rt+8}", f"Invoice sub-total before VAT: {REC['invoice_subtotal']:,}"),
       ("Difference", f"=B{rt+9}-{REC['invoice_subtotal']}", "should be zero (rounding)")]
for i, (a, b, c_) in enumerate(rec, start=rt + 5):
    ws.cell(row=i, column=1, value=a).font = f_b; ws.cell(row=i, column=2, value=b).number_format = USD2; ws.cell(row=i, column=3, value=c_).font = f_note
ws.freeze_panes = "E3"
PRJ = "'Projects baseline'"; R_INSCOPE = rt + 1; R_EXCL = rt + 2
def prj(col, row): return f"{PRJ}!${L(col)}${row}"
proj_row = {p["id"]: r0 + i for i, p in enumerate(PROJECTS)}

# ================================================================ Lines (in-scope run-rates)
ws = wb.create_sheet("Lines")
hdr(ws, 1, ["Cost line", "Category", "Run-rate USD/month (baseline, in scope)", "Volume elasticity", "Commitment flag", "Cash cadence", "Source / note"], [52, 24, 22, 12, 44, 34, 100])
LINES = []
for k, pname in enumerate(PRODUCTS):
    LINES.append(dict(name=pname, cat="GCP native", rr=f"={prj(5+k, R_INSCOPE)}", el=ELAST[pname], flag="Variable", cash="Monthly, reseller invoice", note="Sum of in-scope projects, baseline invoice"))
for k, it in enumerate(ITEMS):
    LINES.append(dict(name=it["name"], cat="Allocated account-level", rr=f"={prj(COL_ITEM0+k, R_INSCOPE)}", el=it.get("elasticity", 0), flag=it["flag"], cash=it["cash"], note=it.get("note", "") + "; in-scope share of the allocation"))
AWS_LINE = "AWS usage accounts"; AWS_TOTAL = sum(a["baseline"] for a in AWS)
LINES.append(dict(name=AWS_LINE, cat="AWS", rr=AWS_TOTAL, el=CFG.get("aws_elasticity", 0.3), flag="Variable", cash="Monthly, reseller AWS invoice", note="Baseline invoice by linked account; see Breakdown"))
for t in TERM:
    LINES.append(dict(name=t["name"], cat=t.get("category", "AWS"), rr=t["monthly"], el=0, flag="Committed (ends per Inputs B15)", cash="Monthly, reseller invoice", note=t.get("note", "")))
for i, ln in enumerate(LINES, start=2):
    ws.cell(row=i, column=1, value=ln["name"]).alignment = wrap; ws.cell(row=i, column=2, value=ln["cat"])
    c = ws.cell(row=i, column=3, value=ln["rr"]); c.number_format = USD
    if not isinstance(ln["rr"], str): blue(c)
    blue(ws.cell(row=i, column=4, value=ln["el"]))
    ws.cell(row=i, column=5, value=ln["flag"]); ws.cell(row=i, column=6, value=ln["cash"]); ws.cell(row=i, column=7, value=ln["note"]).font = f_note
    for c in (1, 5, 6, 7): ws.cell(row=i, column=c).alignment = wrap
    if ln["flag"].startswith(("Committed", "Floor")):
        for c in (1, 2, 5, 6): ws.cell(row=i, column=c).fill = fill_com
line_row = {ln["name"]: i for i, ln in enumerate(LINES, start=2)}
def committed(flag): return flag.startswith(("Committed", "Floor"))
FLOOR_LINE = FLOOR_ITEM["name"] if FLOOR_ITEM else None; SUPPORT_LINE = SUPPORT_ITEM["name"] if SUPPORT_ITEM else None
TERM_LINES = [t["name"] for t in TERM]

# ================================================================ Initiatives
ws = wb.create_sheet("Initiatives")
hdr(ws, 1, ["Initiative", "Target cost line", "Scope (projects)", "Saving share P", "Saving share R", "Saving share O", "Added cost P", "Added cost R", "Added cost O", "Start P (t)", "Start R", "Start O", "Ramp months", "Owner", "Assumption / risk"], [54, 34, 10, 10, 10, 10, 10, 10, 10, 9, 9, 9, 9, 26, 70])
for i, it in enumerate(INIT, start=2):
    row = [it["name"], it["target"], it.get("scope", "all"), it["saving"]["P"], it["saving"]["R"], it["saving"]["O"], it.get("added", {}).get("P", 0), it.get("added", {}).get("R", 0), it.get("added", {}).get("O", 0), it["start"]["P"], it["start"]["R"], it["start"]["O"], it["ramp"], it.get("owner", ""), it.get("note", "")]
    for c, v in enumerate(row, start=1):
        cell = ws.cell(row=i, column=c, value=v); cell.alignment = wrap
        if 4 <= c <= 13: blue(cell, PCT if c <= 9 else None)
ws.row_dimensions[1].height = 45
SAVCOL = {"P": "D", "R": "E", "O": "F"}; ADDCOL = {"P": "G", "R": "H", "O": "I"}; STCOL = {"P": "J", "R": "K", "O": "L"}; RAMPCOL = "M"

# ================================================================ Innovation envelope
ENVW = None
if ENV:
    ws = wb.create_sheet("Innovation envelope")
    ws["A1"] = "Discretionary innovation envelope, company level. Amount on Inputs B18, utilization per scenario on the Inputs scenario table, reseller-billed share on Inputs B19. Weights phase the plan year and must sum to 100%."; ws["A1"].font = f_note
    hdr(ws, 2, ["Month", "Phasing weight", "USD at plan (FX on Inputs B4)", "Note"], [12, 14, 26, 90])
    weights = ENV.get("weights") or [1 / len(PLAN_IDX)] * len(PLAN_IDX)
    for k, j in enumerate(PLAN_IDX):
        rr = 3 + k; ws.cell(row=rr, column=1, value=MONTHS[j]).font = f_b
        blue(ws.cell(row=rr, column=2, value=weights[k]), '0.00%'); ws.cell(row=rr, column=3, value=f"={ENVAMT}*{FX}*B{rr}").number_format = USD
    rt_ = 3 + len(PLAN_IDX); ws.cell(row=rt_, column=1, value="Total").font = f_bold
    ws.cell(row=rt_, column=2, value=f"=SUM(B3:B{rt_-1})").number_format = '0.00%'; ws.cell(row=rt_, column=3, value=f"=SUM(C3:C{rt_-1})").number_format = USD
    for i, t in enumerate(ENV.get("governance", []), start=rt_ + 2):
        c = ws.cell(row=i, column=1, value=t); c.font = f_b; c.alignment = wrap; ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4); ws.row_dimensions[i].height = 30
    ENVW = lambda j: f"'Innovation envelope'!$B${3 + PLAN_IDX.index(j)}"

# ================================================================ Shares (allocation keys)
ws = wb.create_sheet("Shares")
ws["A1"] = "Allocation keys from the baseline matrix: how each product line, allocated item, initiative saving and the credit spread over the in-scope projects. Each row sums to 100%."; ws["A1"].font = f_note
hdr(ws, 2, ["Component"] + [p["id"] for p in INSCOPE] + ["Check"], [56] + [13] * len(INSCOPE) + [8])
r = 3
for k, pname in enumerate(PRODUCTS):
    ws.cell(row=r, column=1, value=f"Product: {pname}").font = f_b
    for j, p in enumerate(INSCOPE):
        ws.cell(row=r, column=2 + j, value=f"=IF({prj(5+k, R_INSCOPE)}=0,0,{prj(5+k, proj_row[p['id']])}/{prj(5+k, R_INSCOPE)})").number_format = PCT
    r += 1
for it in ITEMS:
    ws.cell(row=r, column=1, value=f"Allocated item: {it['name']}").font = f_b
    for j, p in enumerate(INSCOPE):
        ws.cell(row=r, column=2 + j, value=f"={prj(COL_TOTAL, proj_row[p['id']])}/{prj(COL_TOTAL, R_INSCOPE)}").number_format = PCT
    r += 1
for it in INIT:
    ws.cell(row=r, column=1, value=f"Initiative: {it['name']}").font = f_b
    if it["target"] not in PRODUCTS:   # savings on an allocated item follow usage share
        for j, p in enumerate(INSCOPE):
            ws.cell(row=r, column=2 + j, value=f"={prj(COL_TOTAL, proj_row[p['id']])}/{prj(COL_TOTAL, R_INSCOPE)}").number_format = PCT
    else:
        pcol = 5 + PRODUCTS.index(it["target"]); scope = it.get("scope", "all")
        ok = (lambda p: True) if scope == "all" else (lambda p: p["env"] in scope.split("+"))
        denom = "+".join(prj(pcol, proj_row[p["id"]]) for p in INSCOPE if ok(p)) or "0"
        for j, p in enumerate(INSCOPE):
            ws.cell(row=r, column=2 + j, value=(f"=IF(({denom})=0,0,{prj(pcol, proj_row[p['id']])}/({denom}))" if ok(p) else "=0")).number_format = PCT
    r += 1
ws.cell(row=r, column=1, value="Credit, by usage share").font = f_b
for j, p in enumerate(INSCOPE):
    ws.cell(row=r, column=2 + j, value=f"={prj(COL_TOTAL, proj_row[p['id']])}/{prj(COL_TOTAL, R_INSCOPE)}").number_format = PCT
SH_FIRST, SH_LAST = 3, r
for rr in range(SH_FIRST, SH_LAST + 1):
    c = ws.cell(row=rr, column=2 + len(INSCOPE), value=f"=SUM({L(2)}{rr}:{L(1+len(INSCOPE))}{rr})"); c.number_format = PCT; c.font = f_note
ws.freeze_panes = "B3"

# ================================================================ Commitments register + cash events
ws = wb.create_sheet("Commitments")
ws["A1"] = "Commitment register (what is locked, for how long, how it spreads month over month) and the cash events by posting month. Events are whole-account amounts."; ws["A1"].font = f_note
hdr(ws, 2, ["Vendor / instrument", "Term", "Total commitment (USD)", "Cash cadence", "MoM spread (USD/month)", "Next cash event", "Renewal / decision point", "Flag / note"], [40, 30, 18, 26, 18, 26, 26, 80])
for i, row in enumerate(CFG.get("commitments_register", []), start=3):
    for c, key in enumerate(["vendor", "term", "total", "cadence", "mom_spread", "next_event", "decision_point", "note"], start=1):
        cell = ws.cell(row=i, column=c, value=row.get(key)); cell.alignment = wrap; cell.font = f_b
    ws.cell(row=i, column=3).number_format = USD; ws.cell(row=i, column=5).number_format = USD
rE = len(CFG.get("commitments_register", [])) + 5
hdr(ws, rE, ["Cash item (whole account)", "Amount per event", "Posting months"] + MONTHS)
ev_rows = {}
for i, ev in enumerate(CFG["cash_events"], start=rE + 1):
    ws.cell(row=i, column=1, value=ev["name"]).alignment = wrap
    blue(ws.cell(row=i, column=2, value=ev["amount"]), USD); blue(ws.cell(row=i, column=3, value=", ".join(ev["months"])))
    cond = ev.get("switch_equals")
    for j in range(N):
        col = L(4 + j); base = f'IF(ISNUMBER(SEARCH({col}${rE},$C{i})),$B{i},0)'
        if cond is not None: base = f"IF({SWITCH}={cond},{base},0)"
        ws.cell(row=i, column=4 + j, value=f"={base}").number_format = USD
    ev_rows[ev["name"]] = i
r_ev_total = rE + 1 + len(CFG["cash_events"]); ws.cell(row=r_ev_total, column=1, value="Commitment cash events total (whole account)").font = f_bold
for j in range(N):
    col = L(4 + j); c = ws.cell(row=r_ev_total, column=4 + j, value=f"=SUM({col}{rE+1}:{col}{r_ev_total-1})"); c.number_format = USD; c.font = f_bold; c.fill = fill_sub
for j in range(N): ws.column_dimensions[L(4 + j)].width = 12
EVTOT = f"Commitments!$D${r_ev_total}:${L(3+N)}${r_ev_total}"
AUG_EV = sum(ev.get("fallback_in_lag_months", 0) for ev in CFG["cash_events"])

# ================================================================ Projection per scenario
summary_refs, comp_refs, usage_rows_by_key = {}, {}, {}
for key, name in SC:
    ws = wb.create_sheet(f"Projection {key}")
    ws["A1"] = f"{name}: USD per month, amortized, in-scope projects only (allocated items at the in-scope share). Orange = committed."; ws["A1"].font = f_bold
    for j, mth in enumerate(MONTHS):
        c = ws.cell(row=1, column=3 + j, value=mth); c.font = f_h; c.fill = fill_h
    ws.cell(row=2, column=1, value="Month index t").font = f_note
    for j in range(N): ws.cell(row=2, column=3 + j, value=j + 1).font = f_note
    r = 3; ws.cell(row=r, column=1, value="BASELINE").font = f_bold; r += 1
    base_rows = {}
    for ln in LINES:
        n = ln["name"]; lr = line_row[n]
        ws.cell(row=r, column=1, value=n).font = f_b; ws.cell(row=r, column=2, value=("Committed" if committed(ln["flag"]) else "Variable")).font = f_note
        if committed(ln["flag"]): ws.cell(row=r, column=1).fill = fill_com; ws.cell(row=r, column=2).fill = fill_com
        for j in range(N):
            col = L(3 + j); t = f"{col}$2"
            if n == SUPPORT_LINE:
                f = "=" + SUPP + "*(" + "+".join(f"{col}{base_rows[x['name']]}" for x in LINES if x["cat"] == "GCP native") + ")"
            elif n == FLOOR_LINE:
                f = f"=MAX({FMIN}*(1+{FUP})*{prj(COL_SHARE, R_INSCOPE)},Lines!$C${lr}*(1+Lines!$D${lr}*((1+{GROW[key]})^{t}-1)))"
            elif n in TERM_LINES:
                f = f"=IF({t}<={TERMEND},Lines!$C${lr},0)"
            else:
                f = f"=Lines!$C${lr}*(1+Lines!$D${lr}*((1+{GROW[key]})^{t}-1))"
            ws.cell(row=r, column=3 + j, value=f).number_format = USD
        base_rows[n] = r; r += 1
    r_env = None
    if ENV:
        r_env = r; ws.cell(row=r, column=1, value="Innovation envelope (company level, discretionary)").font = f_b; ws.cell(row=r, column=2, value="Discretionary").font = f_note
        ws.cell(row=r, column=1).fill = fill_in; ws.cell(row=r, column=2).fill = fill_in
        for j in range(N):
            ws.cell(row=r, column=3 + j, value=(f"={ENVAMT}*{FX}*{ENVUTIL[key]}*{ENVW(j)}" if j in PLAN_IDX else "=0")).number_format = USD
        r += 1
    r_bt = r; ws.cell(row=r, column=1, value="Baseline total (incl. envelope)").font = f_bold
    for j in range(N):
        col = L(3 + j); c = ws.cell(row=r, column=3 + j, value=f"=SUM({col}4:{col}{r-1})"); c.number_format = USD; c.font = f_bold; c.fill = fill_sub
    com_rows = [base_rows[x["name"]] for x in LINES if committed(x["flag"])]; var_rows = [base_rows[x["name"]] for x in LINES if not committed(x["flag"])]
    r += 1; r_com = r; ws.cell(row=r, column=1, value="of which committed (MoM spread of contracts and floors)").font = f_b
    for j in range(N):
        col = L(3 + j); c = ws.cell(row=r, column=3 + j, value="=" + ("+".join(f"{col}{x}" for x in com_rows) or "0")); c.number_format = USD; c.fill = fill_com
    r += 1; r_envsub = r; ws.cell(row=r, column=1, value="of which discretionary envelope").font = f_b
    for j in range(N):
        ws.cell(row=r, column=3 + j, value=(f"={L(3+j)}{r_env}" if r_env else "=0")).number_format = USD
    r += 1; r_var = r; ws.cell(row=r, column=1, value="of which variable (usage-billed)").font = f_b
    for j in range(N):
        col = L(3 + j); ws.cell(row=r, column=3 + j, value="=" + "+".join(f"{col}{x}" for x in var_rows)).number_format = USD
    r += 1; r_cashb = r; ws.cell(row=r, column=1, value="of which billed monthly through the reseller (cash basis: variable + terminating commitments)").font = f_note
    for j in range(N):
        col = L(3 + j); ws.cell(row=r, column=3 + j, value=f"={col}{r_var}" + "".join(f"+{col}{base_rows[t]}" for t in TERM_LINES)).number_format = USD
    r += 2; ws.cell(row=r, column=1, value="INITIATIVE SAVINGS (net of added cost, ramped)").font = f_bold; r += 1
    init_rows = {}
    for i, it in enumerate(INIT, start=2):
        ws.cell(row=r, column=1, value=it["name"]).font = f_b; ws.cell(row=r, column=1).alignment = wrap; ws.cell(row=r, column=2, value=it["target"]).font = f_note
        tr = base_rows[it["target"]]
        for j in range(N):
            col = L(3 + j); tt = f"{col}$2"
            ramp = f"MIN(1,MAX(0,({tt}-Initiatives!${STCOL[key]}${i}+1)/Initiatives!${RAMPCOL}${i}))"
            ws.cell(row=r, column=3 + j, value=f"={col}{tr}*(Initiatives!${SAVCOL[key]}${i}-Initiatives!${ADDCOL[key]}${i})*{ramp}").number_format = USD
        init_rows[i] = r; r += 1
    r_st = r; ws.cell(row=r, column=1, value="Initiative savings total").font = f_bold
    for j in range(N):
        col = L(3 + j); c = ws.cell(row=r, column=3 + j, value=(f"=SUM({col}{min(init_rows.values())}:{col}{max(init_rows.values())})" if init_rows else "=0")); c.number_format = USD; c.font = f_bold; c.fill = fill_sub
    r += 1; r_stv = r; ws.cell(row=r, column=1, value="of which on variable lines (cash-relevant within contract terms)").font = f_note
    var_init = [init_rows[i] for i, it in enumerate(INIT, start=2) if it["target"] != FLOOR_LINE]
    for j in range(N):
        col = L(3 + j); ws.cell(row=r, column=3 + j, value="=" + ("+".join(f"{col}{x}" for x in var_init) or "0")).number_format = USD
    r += 2; r_cr = r; ws.cell(row=r, column=1, value="Milestone credit, in-scope share").font = f_b
    for j in range(N):
        col = L(3 + j); ws.cell(row=r, column=3 + j, value=f"=IF({col}$2={CREDM[key]},{CREDA[key]}*{prj(COL_SHARE, R_INSCOPE)},0)").number_format = USD
    r += 1; r_net = r; ws.cell(row=r, column=1, value="NET CLOUD COST (USD, amortized, in scope)").font = f_bold
    for j in range(N):
        col = L(3 + j); c = ws.cell(row=r, column=3 + j, value=f"={col}{r_bt}-{col}{r_st}-{col}{r_cr}"); c.number_format = USD; c.font = f_bold; c.fill = fill_sub
    r += 1; r_neteur = r; ws.cell(row=r, column=1, value=f"Net cloud cost ({CCY})").font = f_bold
    for j in range(N):
        col = L(3 + j); ws.cell(row=r, column=3 + j, value=f"={col}{r_net}/{FX}").number_format = USD
    r += 2; ws.cell(row=r, column=1, value="REVENUE AND CAP").font = f_bold; r += 1
    r_rev = r; ws.cell(row=r, column=1, value=f"Revenue ({CCY}, plan phased flat)").font = f_b
    for j in range(N): ws.cell(row=r, column=3 + j, value=f"={REV_PLAN if j in PLAN_IDX else REV_CUR}/12").number_format = USD
    r += 1; r_cap = r; ws.cell(row=r, column=1, value=f"Cloud cap at KPI share ({CCY})").font = f_b
    for j in range(N):
        col = L(3 + j); ws.cell(row=r, column=3 + j, value=f"={col}{r_rev}*{CAP}").number_format = USD
    r += 1; ws.cell(row=r, column=1, value="Net cloud cost / revenue").font = f_bold
    for j in range(N):
        col = L(3 + j); c = ws.cell(row=r, column=3 + j, value=f"={col}{r_neteur}/{col}{r_rev}"); c.number_format = PCT; c.font = f_bold
    r += 1; ws.cell(row=r, column=1, value="Net cloud cost / revenue, excluding the envelope").font = f_b
    for j in range(N):
        col = L(3 + j); ws.cell(row=r, column=3 + j, value=f"=({col}{r_net}-{col}{r_envsub})/{FX}/{col}{r_rev}").number_format = PCT
    r += 1; ws.cell(row=r, column=1, value=f"Headroom vs cap ({CCY}, positive = under; envelope counted per Inputs B20)").font = f_bold
    for j in range(N):
        col = L(3 + j); c = ws.cell(row=r, column=3 + j, value=f"={col}{r_cap}-({col}{r_net}-(1-{ENVKPI})*{col}{r_envsub})/{FX}"); c.number_format = USD; c.font = f_bold
    r += 2; ws.cell(row=r, column=1, value="COMPONENTS for the project breakdown (products, allocated items, initiative savings as negatives, credit as negative)").font = f_note; r += 1
    comp_first = r
    for pname in PRODUCTS:
        ws.cell(row=r, column=1, value=pname).font = f_note
        for j in range(N): ws.cell(row=r, column=3 + j, value=f"={L(3+j)}{base_rows[pname]}").number_format = USD
        r += 1
    for it in ITEMS:
        ws.cell(row=r, column=1, value=it["name"]).font = f_note
        for j in range(N): ws.cell(row=r, column=3 + j, value=f"={L(3+j)}{base_rows[it['name']]}").number_format = USD
        r += 1
    for i, it in enumerate(INIT, start=2):
        ws.cell(row=r, column=1, value=f"saving: {it['name']}").font = f_note
        for j in range(N): ws.cell(row=r, column=3 + j, value=f"=-{L(3+j)}{init_rows[i]}").number_format = USD
        r += 1
    ws.cell(row=r, column=1, value="credit (negative)").font = f_note
    for j in range(N): ws.cell(row=r, column=3 + j, value=f"=-{L(3+j)}{r_cr}").number_format = USD
    comp_last = r; r += 2
    ws.cell(row=r, column=1, value="WHOLE-ACCOUNT HELPERS for the vendor cash flow (excluded projects at baseline usage with the same growth, no savings)").font = f_note; r += 1
    excl_rows = []
    for k, pname in enumerate(PRODUCTS):
        ws.cell(row=r, column=1, value=f"excluded projects: {pname}").font = f_note
        for j in range(N):
            col = L(3 + j); t = f"{col}$2"
            ws.cell(row=r, column=3 + j, value=f"={prj(5+k, R_EXCL)}*(1+Lines!$D${line_row[pname]}*((1+{GROW[key]})^{t}-1))").number_format = USD
        excl_rows.append(r); r += 1
    r_gcp_whole = r; ws.cell(row=r, column=1, value="GCP usage, whole account, net of in-scope initiative savings on product lines").font = f_b
    gcp_init = [init_rows[i] for i, it in enumerate(INIT, start=2) if it["target"] in PRODUCTS]
    for j in range(N):
        col = L(3 + j)
        f = "=" + "+".join(f"{col}{base_rows[p]}" for p in PRODUCTS) + "+" + "+".join(f"{col}{x}" for x in excl_rows) + ("-(" + "+".join(f"{col}{x}" for x in gcp_init) + ")" if gcp_init else "")
        ws.cell(row=r, column=3 + j, value=f).number_format = USD
    r += 1
    whole_items = {}
    for it in ITEMS:
        if it.get("usage_billed"):
            ws.cell(row=r, column=1, value=f"{it['name']}, whole account").font = f_b
            for j in range(N):
                col = L(3 + j); ws.cell(row=r, column=3 + j, value=f"={col}{base_rows[it['name']]}/{prj(COL_SHARE, R_INSCOPE)}").number_format = USD
            whole_items[it["name"]] = r; r += 1
    r += 1; ws.cell(row=r, column=1, value=f"{PLAN} totals").font = f_bold; r += 1
    f27, l27 = L(3 + P0), L(2 + N); tot_rows = {}
    for label, rr_ in ((f"Baseline {PLAN} (USD)", r_bt), (f"Committed {PLAN} (USD)", r_com), (f"Envelope {PLAN} (USD)", r_envsub), (f"Initiative savings {PLAN} (USD)", r_st), (f"Credits {PLAN} (USD)", r_cr), (f"Net {PLAN} (USD)", r_net), (f"Net {PLAN} ({CCY})", r_neteur), (f"Cap {PLAN} ({CCY})", r_cap), (f"Revenue {PLAN} ({CCY})", r_rev)):
        ws.cell(row=r, column=1, value=label).font = f_b
        c = ws.cell(row=r, column=3, value=f"=SUM({f27}{rr_}:{l27}{rr_})"); c.number_format = USD; c.font = f_bold; tot_rows[label] = r; r += 1
    ws.cell(row=r, column=1, value=f"Net {PLAN} / revenue (envelope included)").font = f_bold
    c = ws.cell(row=r, column=3, value=f"=C{tot_rows[f'Net {PLAN} ({CCY})']}/C{tot_rows[f'Revenue {PLAN} ({CCY})']}"); c.number_format = PCT; c.font = f_bold; tot_rows["ratio"] = r; r += 1
    ws.cell(row=r, column=1, value=f"Net {PLAN} / revenue, excluding the envelope").font = f_b
    c = ws.cell(row=r, column=3, value=f"=(C{tot_rows[f'Net {PLAN} (USD)']}-C{tot_rows[f'Envelope {PLAN} (USD)']})/{FX}/C{tot_rows[f'Revenue {PLAN} ({CCY})']}"); c.number_format = PCT; tot_rows["ratio_x"] = r; r += 1
    ws.cell(row=r, column=1, value=f"Gap to cap {PLAN} ({CCY}, negative = over; envelope per Inputs B20)").font = f_bold
    c = ws.cell(row=r, column=3, value=f"=C{tot_rows[f'Cap {PLAN} ({CCY})']}-(C{tot_rows[f'Net {PLAN} (USD)']}-(1-{ENVKPI})*C{tot_rows[f'Envelope {PLAN} (USD)']})/{FX}"); c.number_format = USD; c.font = f_bold; tot_rows["gap"] = r
    ws.column_dimensions["A"].width = 60; ws.column_dimensions["B"].width = 12
    for j in range(N): ws.column_dimensions[L(3 + j)].width = 12
    ws.freeze_panes = "C3"
    summary_refs[key] = dict(sheet=f"Projection {key}", tot=tot_rows, r_cashb=r_cashb, r_stv=r_stv, r_cr=r_cr, r_net=r_net, r_gcp_whole=r_gcp_whole, whole_items=whole_items, r_aws=base_rows[AWS_LINE], r_term={t: base_rows[t] for t in TERM_LINES}, r_env=r_env)
    comp_refs[key] = (comp_first, comp_last)

# ================================================================ Breakdown per scenario (projects + AWS accounts, amortized)
GROUP_ORDER = CFG.get("group_order") or sorted({p["group"] for p in INSCOPE})
NI = len(INIT); NIT = len(ITEMS)
for key, name in SC:
    ws = wb.create_sheet(f"Breakdown {key}"); P = f"'Projection {key}'"; cf_, cl_ = comp_refs[key]; R = summary_refs[key]
    ws["A1"] = f"{name}: net cloud cost by GCP project and AWS account (USD per month, amortized). Project = its share of every product line and allocated item, minus its share of the initiatives in scope, minus its share of the credit."; ws["A1"].font = f_bold
    for j, mth in enumerate(MONTHS):
        c = ws.cell(row=1, column=3 + j, value=mth); c.font = f_h; c.fill = fill_h
    for c, v in ((1, "Project / account"), (2, "Group"), (3 + N, f"{PLAN} total")):
        ws.cell(row=1, column=c, value=v).font = f_h; ws.cell(row=1, column=c).fill = fill_h
    r = 2; group_rows = {}
    for g in GROUP_ORDER:
        members = [p for p in INSCOPE if p["group"] == g]
        if not members: continue
        first = r
        for p in members:
            scol = L(2 + INSCOPE.index(p))
            ws.cell(row=r, column=1, value=p["id"]).font = f_b; ws.cell(row=r, column=2, value=g).font = f_note
            for j in range(N):
                col = L(3 + j); ws.cell(row=r, column=3 + j, value=f"=SUMPRODUCT({P}!{col}${cf_}:{col}${cl_},Shares!${scol}${SH_FIRST}:${scol}${SH_LAST})").number_format = USD
            r += 1
        ws.cell(row=r, column=1, value=f"Group total: {g}").font = f_bold; ws.cell(row=r, column=1).fill = fill_sub
        for j in range(N):
            col = L(3 + j); c = ws.cell(row=r, column=3 + j, value=f"=SUM({col}{first}:{col}{r-1})"); c.number_format = USD; c.font = f_bold; c.fill = fill_sub
        group_rows[g] = r; r += 1
    r_envb = None
    if ENV:
        r_envb = r; ws.cell(row=r, column=1, value="Innovation envelope (company level, not allocated)").font = f_bold; ws.cell(row=r, column=1).fill = fill_in
        for j in range(N):
            c = ws.cell(row=r, column=3 + j, value=f"={P}!{L(3+j)}{R['r_env']}"); c.number_format = USD; c.fill = fill_in
        r += 1
    r_gcp = r; ws.cell(row=r, column=1, value="GCP in-scope total").font = f_bold
    for j in range(N):
        col = L(3 + j); c = ws.cell(row=r, column=3 + j, value="=" + "+".join(f"{col}{group_rows[g]}" for g in group_rows) + (f"+{col}{r_envb}" if r_envb else "")); c.number_format = USD; c.font = f_bold
    r += 2; ws.cell(row=r, column=1, value="AWS accounts").font = f_bold; r += 1; aws_first = r
    for a in AWS:
        ws.cell(row=r, column=1, value=a["name"]).font = f_b; ws.cell(row=r, column=2, value="aws").font = f_note
        for j in range(N):
            col = L(3 + j); ws.cell(row=r, column=3 + j, value=f"={P}!{col}{R['r_aws']}*{a['baseline']}/{AWS_TOTAL}").number_format = USD
        r += 1
    for t in TERM_LINES:
        ws.cell(row=r, column=1, value=f"{t} (until Inputs B15)").font = f_b; ws.cell(row=r, column=2, value="aws").font = f_note
        for j in range(N): ws.cell(row=r, column=3 + j, value=f"={P}!{L(3+j)}{R['r_term'][t]}").number_format = USD
        r += 1
    r_awst = r; ws.cell(row=r, column=1, value="AWS total").font = f_bold; ws.cell(row=r, column=1).fill = fill_sub
    for j in range(N):
        col = L(3 + j); c = ws.cell(row=r, column=3 + j, value=f"=SUM({col}{aws_first}:{col}{r-1})"); c.number_format = USD; c.font = f_bold; c.fill = fill_sub
    r += 2; r_grand = r; ws.cell(row=r, column=1, value="GRAND TOTAL (USD, amortized)").font = f_bold
    for j in range(N):
        col = L(3 + j); c = ws.cell(row=r, column=3 + j, value=f"={col}{r_gcp}+{col}{r_awst}"); c.number_format = USD; c.font = f_bold; c.fill = fill_sub
    r += 1; ws.cell(row=r, column=1, value="Check vs Projection net (should be zero)").font = f_note
    for j in range(N):
        col = L(3 + j); ws.cell(row=r, column=3 + j, value=f"={col}{r_grand}-{P}!{col}{R['r_net']}").number_format = USD
    for rr_ in range(2, r_grand + 1):
        if ws.cell(row=rr_, column=1).value and ws.cell(row=rr_, column=3).value is not None:
            c = ws.cell(row=rr_, column=3 + N, value=f"=SUM({L(3+P0)}{rr_}:{L(2+N)}{rr_})"); c.number_format = USD; c.font = f_bold
    r += 2; ws.cell(row=r, column=1, value="HELPER: GCP usage net by project (products minus product-line savings; feeds Cash by project)").font = f_note; r += 1
    usage_rows = {}
    floor_idx = [i for i, it in enumerate(INIT) if it["target"] == FLOOR_LINE]
    for p in INSCOPE:
        scol = L(2 + INSCOPE.index(p)); ws.cell(row=r, column=1, value=f"usage net: {p['id']}").font = f_note
        for j in range(N):
            col = L(3 + j)
            f = f"=SUMPRODUCT({P}!{col}${cf_}:{col}${cf_+NPR-1},Shares!${scol}${SH_FIRST}:${scol}${SH_FIRST+NPR-1})"
            if NI: f += f"+SUMPRODUCT({P}!{col}${cf_+NPR+NIT}:{col}${cf_+NPR+NIT+NI-1},Shares!${scol}${SH_FIRST+NPR+NIT}:${scol}${SH_FIRST+NPR+NIT+NI-1})"
            for fi in floor_idx: f += f"-{P}!{col}${cf_+NPR+NIT+fi}*Shares!${scol}${SH_FIRST+NPR+NIT+fi}"
            ws.cell(row=r, column=3 + j, value=f).number_format = USD
        usage_rows[p["id"]] = r; r += 1
    usage_rows_by_key[key] = usage_rows
    ws.column_dimensions["A"].width = 62; ws.column_dimensions["B"].width = 10
    for j in range(N + 1): ws.column_dimensions[L(3 + j)].width = 12
    ws.freeze_panes = "C2"

# ================================================================ Cash flow per scenario, by vendor
cash_refs = {}
FB = CFG.get("cash_fallback", {})
for key, name in SC:
    P = f"'Projection {key}'"; ws = wb.create_sheet(f"Cash flow {key}"); R = summary_refs[key]; first, last_ = L(3), L(2 + N)
    ws["A1"] = f"{name}: cash out by vendor, USD per month before VAT. Usage and marketplace postings for month M are invoiced early M+1 at Net 30, so everything reseller-billed is shifted by Inputs B11. Basis per Inputs B17."; ws["A1"].font = f_bold
    for j, mth in enumerate(MONTHS):
        c = ws.cell(row=1, column=3 + j, value=mth); c.font = f_h; c.fill = fill_h
    for c, v in ((1, "Vendor"), (2, "Timing"), (3 + N, f"{PLAN} total")):
        ws.cell(row=1, column=c, value=v).font = f_h; ws.cell(row=1, column=c).fill = fill_h
    for j in range(N): ws.cell(row=2, column=3 + j, value=j + 1).font = f_note
    scale_scope = f"IF({CASHB}=1,1,{prj(COL_SHARE, R_INSCOPE)})"
    def lagged(src_row, fallback):
        return lambda j: f"IF({L(3+j)}$2-{LAG}<1,{fallback},INDEX({P}!${first}${src_row}:${last_}${src_row},1,{L(3+j)}$2-{LAG}))"
    def lagged_ev(ev_row, fallback=0):
        return lambda j: f"IF({L(3+j)}$2-{LAG}<1,{fallback},INDEX(Commitments!$D${ev_row}:${L(3+N)}${ev_row},1,{L(3+j)}$2-{LAG}))"
    vendors = [("GCP usage (all projects, via the reseller invoice)", "usage month + lag", lagged(R["r_gcp_whole"], FB.get("gcp_usage", REC["project_usage"])), scale_scope),
               ("GCP Support (via the reseller invoice)", "usage month + lag", (lambda j: f"{SUPP}*(" + lagged(R["r_gcp_whole"], FB.get("gcp_usage", REC["project_usage"]))(j) + ")"), scale_scope)]
    for ev in CFG["cash_events"]:
        vendors.append((ev["name"], "posting month + lag", lagged_ev(ev_rows[ev["name"]], ev.get("fallback_in_lag_months", 0)), scale_scope))
    for it in ITEMS:
        if it.get("usage_billed"):
            vendors.append((f"{it['name']}: monthly usage", "usage month + lag", lagged(R["whole_items"][it["name"]], it.get("aug_invoice", 0)), scale_scope))
    vendors.append(("AWS usage accounts (via the reseller AWS invoice)", "usage month + lag", lagged(R["r_aws"], FB.get("aws_usage", AWS_TOTAL)), "1"))
    for t in TERM:
        vendors.append((f"{t['name']} (until Inputs B15)", "usage month + lag", lagged(R["r_term"][t["name"]], t["monthly"]), "1"))
    if ENV:
        vendors.append(("Innovation envelope: reseller-billed capacity", "usage month + lag", lagged(R["r_env"], 0), f"{ENVRES}"))
        vendors.append(("Innovation envelope: direct-billed AI and SaaS", "usage month + 1", (lambda j: f"IF({L(3+j)}$2-1<1,0,INDEX({P}!${first}${R['r_env']}:${last_}${R['r_env']},1,{L(3+j)}$2-1))"), f"(1-{ENVRES})"))
    for extra in CFG.get("cash_placeholders", []):
        vendors.append((extra, "TBC", lambda j: "0", "1"))
    vendors.append(("Milestone credit (reduces the invoice it posts on)", "posting month + lag", (lambda j: f"-IF({L(3+j)}$2-{LAG}<1,0,IF(INDEX({P}!$C$2:${last_}$2,1,{L(3+j)}$2-{LAG})={CREDM[key]},{CREDA[key]},0))"), scale_scope))
    r = 3; vrows = []
    for vn, timing, fmaker, scale in vendors:
        ws.cell(row=r, column=1, value=vn).font = f_b; ws.cell(row=r, column=1).alignment = wrap; ws.cell(row=r, column=2, value=timing).font = f_note
        for j in range(N): ws.cell(row=r, column=3 + j, value=f"=({fmaker(j)})*{scale}").number_format = USD
        if timing.startswith("posting"):
            for c in (1, 2): ws.cell(row=r, column=c).fill = fill_com
        vrows.append(r); r += 1
    r_tot = r; ws.cell(row=r, column=1, value="TOTAL CASH OUT (USD, before VAT)").font = f_bold
    for j in range(N):
        col = L(3 + j); c = ws.cell(row=r, column=3 + j, value=f"=SUM({col}{vrows[0]}:{col}{vrows[-1]})"); c.number_format = USD; c.font = f_bold; c.fill = fill_sub
    vat = CFG.get("vat_rate", 0)
    r += 1; ws.cell(row=r, column=1, value=f"Total including VAT at {vat:.0%}").font = f_b
    for j in range(N): ws.cell(row=r, column=3 + j, value=f"={L(3+j)}{r_tot}*(1+{vat})").number_format = USD
    r += 1; ws.cell(row=r, column=1, value=f"Total cash out ({CCY}, before VAT)").font = f_b
    for j in range(N): ws.cell(row=r, column=3 + j, value=f"={L(3+j)}{r_tot}/{FX}").number_format = USD
    r += 2; ws.cell(row=r, column=1, value=f"{PLAN} cash out (USD, before VAT)").font = f_bold
    c = ws.cell(row=r, column=3, value=f"=SUM({L(3+P0)}{r_tot}:{last_}{r_tot})"); c.number_format = USD; c.font = f_bold; r_c27 = r
    r += 1; ws.cell(row=r, column=1, value=f"Peak month cash out {PLAN} (USD, before VAT)").font = f_bold
    c = ws.cell(row=r, column=3, value=f"=MAX({L(3+P0)}{r_tot}:{last_}{r_tot})"); c.number_format = USD; c.font = f_bold; r_peak = r
    for rr_ in vrows + [r_tot]:
        c = ws.cell(row=rr_, column=3 + N, value=f"=SUM({L(3+P0)}{rr_}:{L(2+N)}{rr_})"); c.number_format = USD; c.font = f_bold
    r += 2
    for t in CFG.get("cash_notes", []):
        c = ws.cell(row=r, column=1, value=t); c.font = f_note; c.alignment = wrap; ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10); ws.row_dimensions[r].height = 30; r += 1
    ws.column_dimensions["A"].width = 78; ws.column_dimensions["B"].width = 18
    for j in range(N + 1): ws.column_dimensions[L(3 + j)].width = 12
    ws.freeze_panes = "C3"
    cash_refs[key] = dict(sheet=f"Cash flow {key}", r_c27=r_c27, r_peak=r_peak, r_tot=r_tot)

# ================================================================ Cash by project per scenario
for key, name in SC:
    P = f"'Projection {key}'"; B = f"'Breakdown {key}'"; R = summary_refs[key]; U = usage_rows_by_key[key]
    ws = wb.create_sheet(f"Cash by project {key}"); first, last_ = L(3), L(2 + N)
    ws["A1"] = f"{name}: cash out by GCP project and AWS account, USD per month before VAT, whole account. Project = its usage net of its initiatives, plus its share of Support, the vendor payments, usage-billed marketplace items and the credit (share = baseline usage share), all on the Net 30 lag."; ws["A1"].font = f_bold
    for j, mth in enumerate(MONTHS):
        c = ws.cell(row=1, column=3 + j, value=mth); c.font = f_h; c.fill = fill_h
    for c, v in ((1, "Project / account"), (2, "Group"), (3 + N, f"{PLAN} total")):
        ws.cell(row=1, column=c, value=v).font = f_h; ws.cell(row=1, column=c).fill = fill_h
    for j in range(N): ws.cell(row=2, column=3 + j, value=j + 1).font = f_note
    def shared_terms(j, share_ref):
        col = L(3 + j); t = f"{col}$2"
        ev = f"IF({t}-{LAG}<1,{AUG_EV},INDEX({EVTOT},1,{t}-{LAG}))"
        ub = "+".join(f"IF({t}-{LAG}<1,{it.get('aug_invoice', 0)},INDEX({P}!${first}${R['whole_items'][it['name']]}:${last_}${R['whole_items'][it['name']]},1,{t}-{LAG}))" for it in ITEMS if it.get("usage_billed")) or "0"
        cred = f"IF({t}-{LAG}<1,0,IF(INDEX({P}!$C$2:${last_}$2,1,{t}-{LAG})={CREDM[key]},{CREDA[key]},0))"
        return f"({ev}+{ub}-{cred})*{share_ref}"
    r = 3; rows_all = []
    for p in INSCOPE:
        prow = proj_row[p["id"]]; share_ref = prj(COL_SHARE, prow); urow = U[p["id"]]
        ws.cell(row=r, column=1, value=p["id"]).font = f_b; ws.cell(row=r, column=2, value=p["group"]).font = f_note
        for j in range(N):
            col = L(3 + j); t = f"{col}$2"
            usage = f"IF({t}-{LAG}<1,{prj(COL_TOTAL, prow)},INDEX({B}!${first}${urow}:${last_}${urow},1,{t}-{LAG}))"
            ws.cell(row=r, column=3 + j, value=f"=({usage})*(1+{SUPP})+{shared_terms(j, share_ref)}").number_format = USD
        rows_all.append(r); r += 1
    r += 1; ws.cell(row=r, column=1, value="Paid by the company but outside the in-scope projection (excluded projects at baseline usage with growth, no savings, plus their share of the allocated payments)").font = f_bold; r += 1
    for p in PROJECTS:
        if p["in_scope"] == 1: continue
        prow = proj_row[p["id"]]; share_ref = prj(COL_SHARE, prow)
        ws.cell(row=r, column=1, value=p["id"]).font = f_b; ws.cell(row=r, column=2, value=p["group"]).font = f_note
        for c in (1, 2): ws.cell(row=r, column=c).fill = fill_x
        for j in range(N):
            col = L(3 + j); t = f"{col}$2"
            grown = "+".join(f"{prj(5+k, prow)}*(1+Lines!$D${line_row[PRODUCTS[k]]}*((1+{GROW[key]})^({t}-{LAG})-1))" for k in range(NPR))
            usage = f"IF({t}-{LAG}<1,{prj(COL_TOTAL, prow)},{grown})"
            ws.cell(row=r, column=3 + j, value=f"=({usage})*(1+{SUPP})+{shared_terms(j, share_ref)}").number_format = USD
        rows_all.append(r); r += 1
    if ENV:
        r += 1; ws.cell(row=r, column=1, value="Innovation envelope (company level)").font = f_bold; r += 1
        ws.cell(row=r, column=1, value="Envelope: reseller-billed capacity (usage month + lag)").font = f_b
        for j in range(N):
            t = f"{L(3+j)}$2"; ws.cell(row=r, column=3 + j, value=f"=IF({t}-{LAG}<1,0,INDEX({P}!${first}${R['r_env']}:${last_}${R['r_env']},1,{t}-{LAG}))*{ENVRES}").number_format = USD
        rows_all.append(r); r += 1
        ws.cell(row=r, column=1, value="Envelope: direct-billed AI and SaaS (usage month + 1)").font = f_b
        for j in range(N):
            t = f"{L(3+j)}$2"; ws.cell(row=r, column=3 + j, value=f"=IF({t}-1<1,0,INDEX({P}!${first}${R['r_env']}:${last_}${R['r_env']},1,{t}-1))*(1-{ENVRES})").number_format = USD
        rows_all.append(r); r += 1
    r += 1; ws.cell(row=r, column=1, value="AWS accounts (usage month + lag)").font = f_bold; r += 1
    for a in AWS:
        ws.cell(row=r, column=1, value=a["name"]).font = f_b; ws.cell(row=r, column=2, value="aws").font = f_note
        for j in range(N):
            t = f"{L(3+j)}$2"; ws.cell(row=r, column=3 + j, value=f"=IF({t}-{LAG}<1,{a['baseline']},INDEX({P}!${first}${R['r_aws']}:${last_}${R['r_aws']},1,{t}-{LAG})*{a['baseline']}/{AWS_TOTAL})").number_format = USD
        rows_all.append(r); r += 1
    for t_ in TERM:
        ws.cell(row=r, column=1, value=f"{t_['name']} (until Inputs B15)").font = f_b; ws.cell(row=r, column=2, value="aws").font = f_note
        for j in range(N):
            t = f"{L(3+j)}$2"; ws.cell(row=r, column=3 + j, value=f"=IF({t}-{LAG}<1,{t_['monthly']},INDEX({P}!${first}${R['r_term'][t_['name']]}:${last_}${R['r_term'][t_['name']]},1,{t}-{LAG}))").number_format = USD
        rows_all.append(r); r += 1
    r += 1; r_tot = r; ws.cell(row=r, column=1, value="TOTAL CASH OUT (USD, before VAT, whole account)").font = f_bold
    for j in range(N):
        col = L(3 + j); c = ws.cell(row=r, column=3 + j, value="=" + "+".join(f"{col}{x}" for x in rows_all)); c.number_format = USD; c.font = f_bold; c.fill = fill_sub
    r += 1; ws.cell(row=r, column=1, value="Check vs Cash flow vendor total (zero when Inputs B17 = 1 and no cash placeholders)").font = f_note
    for j in range(N):
        col = L(3 + j); ws.cell(row=r, column=3 + j, value=f"={col}{r_tot}-'Cash flow {key}'!{col}{cash_refs[key]['r_tot']}").number_format = USD
    for rr_ in rows_all + [r_tot]:
        c = ws.cell(row=rr_, column=3 + N, value=f"=SUM({L(3+P0)}{rr_}:{L(2+N)}{rr_})"); c.number_format = USD; c.font = f_bold
    ws.column_dimensions["A"].width = 64; ws.column_dimensions["B"].width = 11
    for j in range(N + 1): ws.column_dimensions[L(3 + j)].width = 12
    ws.freeze_panes = "C3"

# ================================================================ Scenario summary
ws = wb.create_sheet("Scenario summary", 0)
ws["A1"] = f"{COMPANY} cloud cost {PLAN}: three scenarios, in scope. USD unless marked {CCY}."; ws["A1"].font = f_t
hdr(ws, 3, ["Measure", "Pessimistic", "Realistic", "Optimistic", "What moves it"], [52, 16, 16, 16, 90])
rows = [(f"Baseline {PLAN} (no initiatives, no credits)", f"Baseline {PLAN} (USD)", "Volume growth x per-line elasticity; account-level items at the in-scope usage share"),
        ("of which committed contracts and floors (MoM spread)", f"Committed {PLAN} (USD)", "Locked regardless of usage until renewal"),
        (f"Innovation envelope {PLAN} (company level)", f"Envelope {PLAN} (USD)", "Inputs B18 x utilization x phasing"),
        (f"Initiative savings {PLAN}", f"Initiative savings {PLAN} (USD)", "Initiatives sheet; scoped savings land only on the projects in scope"),
        (f"Credits {PLAN} (in-scope share)", f"Credits {PLAN} (USD)", "Milestone credit x in-scope usage share"),
        (f"Net {PLAN} (USD, amortized)", f"Net {PLAN} (USD)", ""), (f"Net {PLAN} ({CCY}, amortized)", f"Net {PLAN} ({CCY})", "FX on Inputs B4"),
        (f"Net {PLAN} / revenue (envelope included)", "ratio", "Revenue basis must match the cost scope"),
        (f"Net {PLAN} / revenue, excluding the envelope", "ratio_x", "If innovation is treated as an investment line beside the KPI"),
        (f"Gap to the cap ({CCY}, negative = over)", "gap", "")]
r = 4
for label, k, why in rows:
    ws.cell(row=r, column=1, value=label).font = f_b
    for col, (key, _) in zip((2, 3, 4), SC):
        ref = summary_refs[key]; c = ws.cell(row=r, column=col, value=f"='{ref['sheet']}'!C{ref['tot'][k]}"); c.number_format = PCT if k in ("ratio", "ratio_x") else USD; c.font = f_bold if k in ("ratio", "ratio_x", "gap") else f_b
    ws.cell(row=r, column=5, value=why).font = f_note; ws.cell(row=r, column=5).alignment = wrap; r += 1
for label, k in ((f"{PLAN} cash out (USD, basis per Inputs B17)", "r_c27"), (f"Peak month cash out {PLAN} (USD)", "r_peak")):
    ws.cell(row=r, column=1, value=label).font = f_b
    for col, (key, _) in zip((2, 3, 4), SC):
        ref = cash_refs[key]; c = ws.cell(row=r, column=col, value=f"='{ref['sheet']}'!C{ref[k]}"); c.number_format = USD; c.font = f_bold
    r += 1
r += 1
for t in CFG.get("summary_notes", []):
    c = ws.cell(row=r, column=1, value=t); c.font = f_note; c.alignment = wrap; ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5); ws.row_dimensions[r].height = 40; r += 1

# ================================================================ Sources and gaps
ws = wb.create_sheet("Sources and gaps")
for i, t in enumerate(CFG.get("sources_and_gaps", []), start=1):
    c = ws.cell(row=i, column=1, value=t); c.font = f_bold if t in ("Sources", "Method", "Gaps") else f_b; c.alignment = wrap
ws.column_dimensions["A"].width = 160
wb.save(OUT); print("saved", OUT)

# ================================================================ recalc, tie-outs, CSV
RECALC = CFG.get("recalc_script", "/mnt/skills/public/xlsx/scripts/recalc.py")
if os.path.exists(RECALC):
    res = subprocess.run([sys.executable, RECALC, OUT, "180"], capture_output=True, text=True)
    try:
        summ = json.loads(res.stdout); print("recalc:", summ["status"], "errors", summ["total_errors"], "formulas", summ["total_formulas"])
    except Exception:
        print("recalc output:", res.stdout[-500:], res.stderr[-500:])
    wbv = load_workbook(OUT, data_only=True)
    pb = wbv["Projects baseline"]; diff = pb.cell(row=rt + 10, column=2).value
    print(f"invoice reconciliation difference: {diff}")
    sh = wbv["Shares"]; print("share rows off 100%:", [rr for rr in range(SH_FIRST, SH_LAST + 1) if abs((sh.cell(row=rr, column=2 + len(INSCOPE)).value or 0) - 1) > 1e-6])
    for key, _ in SC:
        bd = wbv[f"Breakdown {key}"]; cp = wbv[f"Cash by project {key}"]
        chk = max(abs(bd.cell(row=rr, column=c).value or 0) for rr in range(2, bd.max_row + 1) if (bd.cell(row=rr, column=1).value or "").startswith("Check") for c in range(3, 3 + N))
        chk2 = max(abs(cp.cell(row=rr, column=c).value or 0) for rr in range(2, cp.max_row + 1) if (cp.cell(row=rr, column=1).value or "").startswith("Check") for c in range(3, 3 + N))
        print(f"{key}: breakdown tie-out max |diff| {chk:.4f}; cash by project tie-out max |diff| {chk2:.4f}")
    ss = wbv["Scenario summary"]
    for rr in range(4, 16):
        v = [ss.cell(row=rr, column=c).value for c in range(1, 5)]
        if v[0] and v[1] is not None: print([round(x, 4) if isinstance(x, float) else x for x in v])
    rows = []
    for key, _ in SC:
        for sheet, kind in ((f"Breakdown {key}", "by_project_account"), (f"Projection {key}", "by_product_line"), (f"Cash flow {key}", "cash_by_vendor"), (f"Cash by project {key}", "cash_by_project_account")):
            w = wbv[sheet]; months = [w.cell(row=1, column=c).value for c in range(3, 3 + N)]
            for rr in range(2, w.max_row + 1):
                a = w.cell(row=rr, column=1).value
                if not a or a.endswith("totals") or a.startswith((f"{PLAN} ", "Peak ", "COMPONENTS", "WHOLE-ACCOUNT", "HELPER", "usage net:", "excluded projects:", "Month index", "Paid by", "Innovation envelope (company", "AWS accounts (")): continue
                vals = [w.cell(row=rr, column=c).value for c in range(3, 3 + N)]
                if all(v is None for v in vals): continue
                for mth, v in zip(months, vals):
                    rows.append([key, kind, mth, a, w.cell(row=rr, column=2).value or "", round(v, 2) if isinstance(v, (int, float)) else v])
    with open(OUT[:-5] + ".csv", "w", newline="") as f:
        wr = csv.writer(f); wr.writerow(["scenario", "view", "month", "line_or_project_or_vendor", "group_or_timing", "value"]); wr.writerows(rows)
    print("csv rows", len(rows))
else:
    print("recalc.py not found: open the workbook in Excel/LibreOffice to compute formulas, then export the CSV")
